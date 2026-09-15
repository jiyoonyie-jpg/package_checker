import base64
import os

def get_file_info(file_bytes: bytes, filename: str) -> dict:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    return {
        "filename": filename,
        "size_kb": round(len(file_bytes) / 1024, 1),
        "extension": ext,
        "is_pdf": ext == "pdf",
        "is_image": ext in ("png", "jpg", "jpeg", "webp"),
        "is_excel": ext in ("xlsx", "xls"),
    }

def extract_excel_text(file_bytes: bytes) -> str:
    """엑셀 시트의 셀 내용을 표 형태 텍스트로 추출합니다."""
    try:
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
        lines = []
        for sheet in wb.worksheets:
            lines.append(f"[시트: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines) if lines else "(빈 엑셀 파일)"
    except Exception as e:
        return f"[엑셀 텍스트 추출 실패: {e}]"

def image_to_base64(file_bytes: bytes, media_type: str = "image/png") -> str:
    return base64.b64encode(file_bytes).decode("utf-8")

def pdf_to_images(file_bytes: bytes) -> list:
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        pages = []
        for i, page in enumerate(doc):
            mat = fitz.Matrix(2.0, 2.0)
            pix = page.get_pixmap(matrix=mat)
            img_bytes = pix.tobytes("png")
            pages.append({
                "page": i + 1,
                "bytes": img_bytes,
                "base64": base64.b64encode(img_bytes).decode("utf-8"),
                "media_type": "image/png",
            })
        doc.close()
        return pages
    except Exception as e:
        return [{"error": str(e)}]

def detect_barcodes(file_bytes: bytes) -> list:
    try:
        from pyzbar.pyzbar import decode
        from PIL import Image
        import io
        img = Image.open(io.BytesIO(file_bytes))
        results = decode(img)
        if not results:
            return [{"error": "not found"}]
        return [{"type": r.type, "data": r.data.decode("utf-8")} for r in results]
    except Exception:
        return [{"error": "barcode lib unavailable"}]

def extract_text_from_pdf(file_bytes: bytes) -> str:
    try:
        import fitz
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        text = ""
        for page in doc:
            text += page.get_text()
        doc.close()
        return text
    except Exception as e:
        return f"[텍스트 추출 실패: {e}]"
