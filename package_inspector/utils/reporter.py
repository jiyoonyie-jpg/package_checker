import io
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

FIELD_NAMES = {
    "product_name": "제품명",
    "manufacturer": "제조자/수입자",
    "content_weight": "내용량",
    "ingredients": "원재료명",
    "nutrition_facts": "영양성분표",
    "expiry_date": "유통기한/소비기한",
    "storage_method": "보관방법",
    "allergen": "알레르기 유발물질",
    "country_of_origin": "원산지",
    "food_type": "식품유형",
    "barcode": "바코드",
}

def _border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

def export_to_excel(result: dict, file_info: dict) -> bytes:
    wb = Workbook()

    # ── Sheet 1: 요약 ──────────────────────────────
    ws = wb.active
    ws.title = "검수 요약"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 45

    header_fill = PatternFill("solid", fgColor="1A1A2E")
    header_font = Font(color="FFFFFF", bold=True, size=13)
    ws.merge_cells("A1:B1")
    ws["A1"] = "패키지 표기사항 AI 검수 결과"
    ws["A1"].fill = header_fill
    ws["A1"].font = header_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    meta = [
        ("파일명", file_info.get("filename", "")),
        ("검수일시", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("종합 점수", f"{result.get('overall_score', 0)} / 100"),
        ("검수 요약", result.get("summary", "")),
    ]
    for r, (k, v) in enumerate(meta, start=2):
        ws[f"A{r}"] = k
        ws[f"B{r}"] = v
        ws[f"A{r}"].font = Font(bold=True)
        ws[f"A{r}"].border = _border()
        ws[f"B{r}"].border = _border()
        ws.row_dimensions[r].height = 18

    # ── Sheet 2: 항목별 결과 ───────────────────────
    ws2 = wb.create_sheet("항목별 결과")
    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 10
    ws2.column_dimensions["C"].width = 30
    ws2.column_dimensions["D"].width = 40

    headers = ["항목", "결과", "확인된 값", "문제점"]
    fill_h = PatternFill("solid", fgColor="0F3460")
    for c, h in enumerate(headers, 1):
        cell = ws2.cell(row=1, column=c, value=h)
        cell.fill = fill_h
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = _border()

    ok_fill = PatternFill("solid", fgColor="E8F5E9")
    fail_fill = PatternFill("solid", fgColor="FCE4EC")
    detected = result.get("detected_fields", {})
    for r, (fid, fname) in enumerate(FIELD_NAMES.items(), start=2):
        fd = detected.get(fid, {})
        found = fd.get("found", False)
        ws2.cell(row=r, column=1, value=fname).border = _border()
        status_cell = ws2.cell(row=r, column=2, value="✅ 있음" if found else "❌ 없음")
        status_cell.fill = ok_fill if found else fail_fill
        status_cell.alignment = Alignment(horizontal="center")
        status_cell.border = _border()
        ws2.cell(row=r, column=3, value=str(fd.get("value") or "")).border = _border()
        ws2.cell(row=r, column=4, value=str(fd.get("issue") or "")).border = _border()
        ws2.row_dimensions[r].height = 18

    # ── Sheet 3: 위반사항 ──────────────────────────
    ws3 = wb.create_sheet("위반 및 주의사항")
    ws3.column_dimensions["A"].width = 12
    ws3.column_dimensions["B"].width = 22
    ws3.column_dimensions["C"].width = 55

    for c, h in enumerate(["구분", "항목", "내용"], 1):
        cell = ws3.cell(row=1, column=c, value=h)
        cell.fill = fill_h
        cell.font = Font(color="FFFFFF", bold=True)
        cell.border = _border()

    r = 2
    sev_fills = {"high": PatternFill("solid", fgColor="FFEBEE"), "medium": PatternFill("solid", fgColor="FFF8E1"), "low": PatternFill("solid", fgColor="F3F9E7")}
    sev_labels = {"high": "🔴 긴급", "medium": "🟡 주의", "low": "🟢 경미"}
    for v in result.get("violations", []):
        sev = v.get("severity", "medium")
        for c, val in enumerate([sev_labels.get(sev, sev), v.get("field", ""), v.get("message", "")], 1):
            cell = ws3.cell(row=r, column=c, value=val)
            cell.fill = sev_fills.get(sev, PatternFill())
            cell.border = _border()
        ws3.row_dimensions[r].height = 18
        r += 1
    for w in result.get("warnings", []):
        for c, val in enumerate(["💡 주의", w.get("field", ""), w.get("message", "")], 1):
            cell = ws3.cell(row=r, column=c, value=val)
            cell.fill = PatternFill("solid", fgColor="E3F2FD")
            cell.border = _border()
        ws3.row_dimensions[r].height = 18
        r += 1

    # ── Sheet 4: 오탈자/맞춤법 ──────────────────────
    typos = result.get("typos", [])
    if typos:
        ws4 = wb.create_sheet("오탈자")
        ws4.column_dimensions["A"].width = 20
        ws4.column_dimensions["B"].width = 25
        ws4.column_dimensions["C"].width = 40
        ws4.column_dimensions["D"].width = 30
        for c, h in enumerate(["위치", "표기된 문구", "문제 설명", "수정 제안"], 1):
            cell = ws4.cell(row=1, column=c, value=h)
            cell.fill = fill_h
            cell.font = Font(color="FFFFFF", bold=True)
            cell.border = _border()
        for r4, t in enumerate(typos, start=2):
            for c, val in enumerate([t.get("location", ""), t.get("found_text", ""),
                                      t.get("issue", ""), t.get("suggestion", "")], 1):
                cell = ws4.cell(row=r4, column=c, value=val)
                cell.fill = PatternFill("solid", fgColor="FFF3E0")
                cell.border = _border()
            ws4.row_dimensions[r4].height = 18

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
